"""Automates the Accenture 'Closed with Insuff' report from the live Bridge database.

Runs the three queries in queries.py (daily Done checks, Advance Tracker,
Antecedent Details) directly against the live MySQL database, saves each result
as a CSV next to this file, and builds the formatted Excel deliverable.
"""
import math
import os
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
import pymysql
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from queries import ADVANCE_QUERY, ANTECEDENT_QUERY, DAILY_QUERY

load_dotenv()

CLIENT_NAME = "Accenture Solutions Private Limited"

DAILY_CSV_NAME = "Yesterday Done checks with other details as per client name - Bridge.csv"
ADVANCE_CSV_NAME = "Advance Tracker.csv"
ANTECEDENT_CSV_NAME = "Antecedent Details.csv"

# Columns the pipeline reads by name. MySQL can report a column in its table-
# defined casing rather than the casing written in the query, so results are
# renamed case-insensitively to these spellings.
DAILY_COLUMNS = [
    "case_check_id", "case_ars_no", "check_status", "case_received_date", "insuff_remarks",
    "Insuff_fulfill_date", "case_flex_field1", "case_flex_field6",
]
ADVANCE_COLUMNS = ["Case_Check_id", "First_Insuff_Date", "verification_source", "Check_unique_name"]
ANTECEDENT_COLUMNS = ["case_check_id", "field_name", "stated_data"]

DOWNLOAD_DIR = Path(__file__).resolve().parent

CHECK_STATUS_FILTER = "Closed with Insufficiency"

# Antecedent Details rows whose stated_data feeds the HRT column.
HRT_FIELD_NAMES = ["Task ID", "Old Task Number", "New Task Number"]

# Default mapping file names looked for next to app.py when no mapping is
# supplied explicitly (e.g. from the Streamlit uploader).
MAPPING_FILE_NAMES = ("mapping.xlsx", "mapping.xls.xlsx")

# Advance Tracker's Check_unique_name -> Component label for column D.
# Used only if no mapping.xlsx (Check_unique_name / Mapped Name columns) is found.
DEFAULT_COMPONENT_MAP = {
    "Date of Birth Verification": "Legal Age to Work",
    "Education Verification W": "Education",
    "India Court Record Database Check": "Criminal Court Check",
    "Previous Employment Verification": "Previous Employment",
    "Global Database Check": "Criminal Database Check",
    "National Identity Check": "Identity Proof Check",
    "Current Employment Company Check": "Current Company",
    "Previous Employment Company Check": "Previous Company",
    "Current Employment Verification": "Current Employment",
    "Professional Reference Check": "Professional Reference Check",
    "Instant Previous Employment Check": "Previous Employment - Tier 3",
    "Instant Current Employment Check": "Current Employment - Tier 3",
    "Education Institution Validation": "Education Institution Validation",
    "India Court Record Check through Law Firm": "India Court Record Check through Law Firm",
}

# (source CSV column, output column header). None = leave blank (filled in by a
# future logic update). "__vendor__" = static "Authbridge" value. The
# "__advance_*__" markers are looked up from the Advance Tracker export and
# "__antecedent_hrt__" from the Antecedent Details export, both matched by
# Case_Check_id / case_check_id.
COLUMN_MAP = [
    ("case_flex_field1", "CID"),
    ("case_flex_field6", "HRC"),
    ("__antecedent_hrt__", "HRT"),
    ("__advance_component__", "Component"),
    ("insuff_remarks", "Insufficiency Details"),
    ("__advance_verification_source__", "University/College/Employment/Company - Name"),
    ("case_received_date", "Date of Initiation"),
    ("__advance_first_insuff_date__", "Insuff Raised date"),
    ("Insuff_fulfill_date", "Insuff Closed date"),
    ("__vendor__", "Vendor Name"),
    ("case_check_id", "case_check_id"),
    ("case_ars_no", "case_ars_no"),
]

# Date columns rendered as dd-mmm-yy (e.g. 09-Jul-26). "Insuff Raised date"
# keeps its full timestamp value underneath the cell - only the display is
# date-only; the raw date and time are still there if you click into it.
DATE_DISPLAY_COLUMNS = {"Date of Initiation", "Insuff Raised date", "Insuff Closed date"}
DATE_NUMBER_FORMAT = "dd-mmm-yy"

# Columns wide enough to hold free-text remarks; wrapped instead of truncated.
WRAP_COLUMNS = {"Insufficiency Details", "University/College/Employment/Company - Name"}

COLUMN_WIDTHS = {
    "CID": 16,
    "HRC": 14,
    "HRT": 16,
    "Component": 24,
    "Insufficiency Details": 55,
    "University/College/Employment/Company - Name": 38,
    "Date of Initiation": 16,
    "Insuff Raised date": 16,
    "Insuff Closed date": 16,
    "Vendor Name": 13,
    "case_check_id": 15,
    "case_ars_no": 14,
}

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BAND_FILL = PatternFill(start_color="F2F6FA", end_color="F2F6FA", fill_type="solid")
GRID_SIDE = Side(style="thin", color="D9D9D9")
GRID_BORDER = Border(left=GRID_SIDE, right=GRID_SIDE, top=GRID_SIDE, bottom=GRID_SIDE)


def compute_date_range():
    """Yesterday's date range, except on Monday where it spans Fri-Sun."""
    today = date.today()
    if today.weekday() == 0:  # Monday
        from_date = today - timedelta(days=3)  # Friday
        to_date = today - timedelta(days=1)  # Sunday
    else:
        from_date = today - timedelta(days=1)
        to_date = from_date
    return f"{from_date} 00:00:00", f"{to_date} 23:59:59"


def connect_db():
    """Open a connection to the live database using DB_* settings from .env."""
    return pymysql.connect(
        host=os.environ["DB_HOST"],
        port=int(os.environ.get("DB_PORT", "3306")),
        user=os.environ["DB_USER"],
        password=os.environ["DB_PASSWORD"],
        database=os.environ.get("DB_NAME", "checkpoint_live"),
        charset="utf8mb4",
        connect_timeout=30,
    )


def fetch_df(conn, query, params=(), expected_columns=()):
    """Run `query` and return the result as a DataFrame, renaming columns
    case-insensitively to the spellings in `expected_columns`."""
    with conn.cursor() as cur:
        cur.execute(query, params)
        columns = [c[0] for c in cur.description]
        rows = cur.fetchall()
    wanted = {name.lower(): name for name in expected_columns}
    columns = [wanted.get(c.lower(), c) for c in columns]
    return pd.DataFrame(rows, columns=columns)


def fetch_in_query(conn, query, values, expected_columns=()):
    """Run a query containing an `{in_list}` slot, bound to `values`."""
    in_list = ",".join(["%s"] * len(values))
    return fetch_df(conn, query.format(in_list=in_list), values, expected_columns)


def save_csv(df, name, log=print):
    path = DOWNLOAD_DIR / name
    df.to_csv(path, index=False)
    log(f"Saved: {path}")
    return path


def load_closed_with_insuff(daily_df):
    return daily_df[daily_df["check_status"] == CHECK_STATUS_FILTER]


def build_hrt_lookup(antecedent_df):
    """case_check_id -> sorted, de-duplicated Task/Old/New Task Number
    stated_data values (comma-joined) for the HRT column."""
    sub = antecedent_df[antecedent_df["field_name"].isin(HRT_FIELD_NAMES)].copy()
    sub["stated_data"] = sub["stated_data"].astype(str).str.strip()
    sub = sub[~sub["stated_data"].isin(["", "nan", "Not Mentioned"])]
    sub["case_check_id"] = sub["case_check_id"].astype(str)
    return sub.groupby("case_check_id")["stated_data"].agg(lambda values: ", ".join(sorted(set(values))))


def load_component_map(source):
    """Read a Check_unique_name -> Mapped Name lookup from an Excel file
    (columns 'Check_unique_name' and 'Mapped Name'). `source` may be a path
    or a file-like object, e.g. a Streamlit file upload."""
    mapping_df = pd.read_excel(source)
    return dict(zip(mapping_df["Check_unique_name"], mapping_df["Mapped Name"]))


def find_default_mapping_file(directory):
    for name in MAPPING_FILE_NAMES:
        candidate = directory / name
        if candidate.exists():
            return candidate
    return None


def estimate_row_height(row, headers):
    """Approximate the row height (points) needed for wrapped text columns,
    since Excel won't auto-fit row height for content it hasn't rendered yet."""
    max_lines = 1
    for header, value in zip(headers, row):
        if header in WRAP_COLUMNS and value:
            chars_per_line = max(int(COLUMN_WIDTHS.get(header, 30) * 1.7), 10)
            max_lines = max(max_lines, math.ceil(len(str(value)) / chars_per_line))
    return min(max_lines * 15, 250)


def write_formatted_report(out, xlsx_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Closed with Insuff"

    headers = list(out.columns)
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = GRID_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, row in enumerate(out.itertuples(index=False), start=2):
        for col_idx, (header, value) in enumerate(zip(headers, row), start=1):
            if pd.isna(value) or value == "":
                value = None
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = GRID_BORDER
            if row_idx % 2 == 0:
                cell.fill = BAND_FILL
            if header in DATE_DISPLAY_COLUMNS and value is not None:
                cell.number_format = DATE_NUMBER_FORMAT
            if header in WRAP_COLUMNS:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row_idx].height = estimate_row_height(row, headers)

    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS.get(header, 16)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(xlsx_path)


def build_insuff_report(closed, advance_df, antecedent_df, out_dir, component_map=None):
    component_map = component_map or DEFAULT_COMPONENT_MAP
    check_ids = closed["case_check_id"].astype(str)

    if advance_df is not None and not advance_df.empty:
        advance_lookup = advance_df.drop_duplicates("Case_Check_id").set_index("Case_Check_id")
        advance_lookup.index = advance_lookup.index.astype(str)
        first_insuff_date = check_ids.map(advance_lookup["First_Insuff_Date"])
        verification_source = check_ids.map(advance_lookup["verification_source"])

        check_unique_name = check_ids.map(advance_lookup["Check_unique_name"])
        component = check_unique_name.map(component_map)
        unmapped = sorted(check_unique_name[check_unique_name.notna() & component.isna()].unique())
        if unmapped:
            print(f"Warning: no Component mapping for Check_unique_name values: {unmapped}")
    else:
        first_insuff_date = ""
        verification_source = ""
        component = ""

    if antecedent_df is not None and not antecedent_df.empty:
        hrt = check_ids.map(build_hrt_lookup(antecedent_df))
    else:
        hrt = ""

    out = pd.DataFrame()
    for source, header in COLUMN_MAP:
        if source is None:
            out[header] = ""
        elif source == "__vendor__":
            out[header] = "Authbridge"
        elif source == "__advance_first_insuff_date__":
            out[header] = first_insuff_date
        elif source == "__advance_verification_source__":
            out[header] = verification_source
        elif source == "__antecedent_hrt__":
            out[header] = hrt
        elif source == "__advance_component__":
            out[header] = component
        else:
            out[header] = closed[source]

    out["Date of Initiation"] = pd.to_datetime(out["Date of Initiation"], errors="coerce")
    out["Insuff Raised date"] = pd.to_datetime(out["Insuff Raised date"], errors="coerce")
    out["Insuff Closed date"] = pd.to_datetime(out["Insuff Closed date"], errors="coerce")

    xlsx_path = out_dir / f"Accenture_Closed_with_Insuff_{date.today()}.xlsx"
    write_formatted_report(out, xlsx_path)
    return xlsx_path


def build_ars_list(closed):
    """Unique ARS numbers (order preserved) bound into the Advance Tracker
    query's `case_ars_no IN (...)` clause."""
    return [str(ars) for ars in dict.fromkeys(closed["case_ars_no"].dropna())]


def build_check_id_list(closed):
    """Unique case_check_id values (order preserved) bound into the Antecedent
    Details query's `case_check_id IN (...)` clause."""
    return list(dict.fromkeys(closed["case_check_id"].dropna().astype(str)))


def run_pipeline(component_map=None, log=print):
    """Run the full query -> filter -> report pipeline and return a dict of
    the produced file paths (plus row counts) keyed by stage name."""
    from_date, to_date = compute_date_range()
    log(f"Querying live database for {from_date} -> {to_date}")

    results = {"from_date": from_date, "to_date": to_date}
    conn = connect_db()
    try:
        daily_df = fetch_df(
            conn, DAILY_QUERY, (from_date, to_date, f"%{CLIENT_NAME}%"), DAILY_COLUMNS
        )
        results["daily_csv"] = save_csv(daily_df, DAILY_CSV_NAME, log)

        closed = load_closed_with_insuff(daily_df)
        results["closed_rows"] = len(closed)
        ars_list = build_ars_list(closed)

        advance_df = None
        if ars_list:
            advance_df = fetch_in_query(conn, ADVANCE_QUERY, ars_list, ADVANCE_COLUMNS)
            results["advance_csv"] = save_csv(advance_df, ADVANCE_CSV_NAME, log)
        else:
            log("No 'Closed with Insufficiency' rows today; skipping Advance Tracker query.")

        check_id_list = build_check_id_list(closed)
        antecedent_df = None
        if check_id_list:
            antecedent_df = fetch_in_query(conn, ANTECEDENT_QUERY, check_id_list, ANTECEDENT_COLUMNS)
            results["antecedent_csv"] = save_csv(antecedent_df, ANTECEDENT_CSV_NAME, log)
        else:
            log("No 'Closed with Insufficiency' rows today; skipping Antecedent Details query.")

        xlsx_path = build_insuff_report(closed, advance_df, antecedent_df, DOWNLOAD_DIR, component_map)
        log(f"Saved: {xlsx_path}")
        results["report_xlsx"] = xlsx_path
    finally:
        conn.close()

    return results


def main():
    mapping_path = find_default_mapping_file(DOWNLOAD_DIR)
    component_map = load_component_map(mapping_path) if mapping_path else DEFAULT_COMPONENT_MAP
    run_pipeline(component_map)


if __name__ == "__main__":
    main()
